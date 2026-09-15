#!/usr/bin/env python3
"""Build the six-slide SIH26237 submission from the official SIH template."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from deck_evidence import load_physical, speaker_notes


NAVY = "0B1F33"
BLUE = "24557A"
TEAL = "15B8A6"
TEAL_DARK = "087E73"
AMBER = "F4B942"
RED = "D95555"
INK = "172A3A"
MUTED = "617386"
PALE = "F3F7F9"
PALE_TEAL = "E8F7F4"
PALE_AMBER = "FFF6DE"
PALE_RED = "FDECEC"
WHITE = "FFFFFF"
GREY = "DCE5EA"
FONT = "Liberation Sans"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def remove_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def clear_slide_body(slide, *, title_slide: bool = False) -> None:
    """Keep official background/footer/logo; remove template body placeholders."""
    for shape in list(slide.shapes):
        text = getattr(shape, "text", "").strip()
        top = shape.top / 914400
        is_logo = shape.shape_type == 13 and shape.left / 914400 > 10.2 and top < 1.3
        is_footer = top > 6.85
        is_background = shape.left == 0 and shape.top == 0 and shape.width >= 12190000
        keep_header = title_slide and text == "SMART INDIA HACKATHON 2026"
        if is_logo or is_footer or is_background or keep_header:
            continue
        remove_shape(shape)


def box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str = WHITE,
    line: str | None = GREY,
    radius: bool = True,
    line_width: float = 1.0,
):
    shape_type = (
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    )
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    return shape


def text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    value: str,
    *,
    size: float = 16,
    color: str = INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.04,
    font: str = FONT,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape


def rich_text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    runs: list[tuple[str, float, str, bool]],
    *,
    margin: float = 0.04,
    valign=MSO_ANCHOR.TOP,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    for value, size, color, bold in runs:
        run = paragraph.add_run()
        run.text = value
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return shape


def bullet_list(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    items: list[str],
    *,
    size: float = 13,
    color: str = INK,
    bullet_color: str = TEAL,
    gap: float = 3,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(gap)
        paragraph.line_spacing = 1.0
        run_bullet = paragraph.add_run()
        run_bullet.text = "●  "
        run_bullet.font.name = FONT
        run_bullet.font.size = Pt(max(size - 2, 8))
        run_bullet.font.color.rgb = rgb(bullet_color)
        run_text = paragraph.add_run()
        run_text.text = item
        run_text.font.name = FONT
        run_text.font.size = Pt(size)
        run_text.font.color.rgb = rgb(color)
    return shape


def pill(slide, x, y, w, label, *, fill=PALE_TEAL, color=TEAL_DARK, size=10.5):
    box(slide, x, y, w, 0.34, fill=fill, line=None)
    text(
        slide,
        x,
        y + 0.01,
        w,
        0.3,
        label,
        size=size,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        margin=0,
    )


def section_header(slide, title_value: str, team_name: str, number: int) -> None:
    text(slide, 0.62, 0.20, 9.55, 0.60, title_value, size=27, color=NAVY, bold=True)
    display_name = "[TEAM NAME]" if team_name.startswith("[") else team_name
    pill(slide, 0.62, 0.83, 1.52, display_name, fill=PALE_TEAL, color=TEAL_DARK, size=9)


def arrow(slide, x: float, y: float, w: float = 0.25, h: float = 0.34, color: str = TEAL):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    return shape


def small_card(slide, x, y, w, h, number, title_value, body, *, accent=TEAL):
    box(slide, x, y, w, h, fill=WHITE, line=GREY)
    circle = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.13), Inches(y + 0.14), Inches(0.34), Inches(0.34)
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = rgb(accent)
    circle.line.fill.background()
    text(
        slide,
        x + 0.13,
        y + 0.14,
        0.34,
        0.34,
        str(number),
        size=10,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        margin=0,
    )
    text(slide, x + 0.55, y + 0.10, w - 0.67, 0.40, title_value, size=11.4, color=NAVY, bold=True)
    text(slide, x + 0.14, y + 0.55, w - 0.28, h - 0.64, body, size=9.4, color=MUTED)


def draw_document_hero(slide) -> None:
    box(slide, 7.95, 1.33, 4.58, 4.92, fill=PALE, line=None)
    doc = box(slide, 9.00, 1.73, 2.18, 2.92, fill=WHITE, line=BLUE, radius=False, line_width=1.6)
    # folded corner
    fold = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RIGHT_TRIANGLE, Inches(10.60), Inches(1.73), Inches(0.58), Inches(0.58)
    )
    fold.rotation = 180
    fold.fill.solid()
    fold.fill.fore_color.rgb = rgb(PALE_TEAL)
    fold.line.color.rgb = rgb(BLUE)
    text(slide, 9.30, 2.30, 1.58, 0.40, "RESTRICTED", size=11.5, color=RED, bold=True, align=PP_ALIGN.CENTER)
    for i, width in enumerate([1.25, 1.48, 1.05, 1.38]):
        bar = box(slide, 9.34, 2.92 + i * 0.28, width, 0.065, fill=GREY, line=None, radius=False)
    # fingerprint-like rings
    for i in range(4):
        ring = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ARC,
            Inches(9.64 - i * 0.08),
            Inches(3.24 - i * 0.08),
            Inches(0.72 + i * 0.16),
            Inches(0.90 + i * 0.16),
        )
        ring.fill.background()
        ring.line.color.rgb = rgb(TEAL)
        ring.line.width = Pt(2.2)
    # ledger validators around document
    positions = [(8.35, 2.00), (11.47, 2.02), (8.38, 4.95), (11.45, 4.92)]
    for index, (x, y) in enumerate(positions, 1):
        node = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(0.54), Inches(0.54))
        node.fill.solid()
        node.fill.fore_color.rgb = rgb(NAVY if index < 4 else MUTED)
        node.line.fill.background()
        text(slide, x, y, 0.54, 0.54, f"V{index}", size=9, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
    pill(slide, 8.64, 5.55, 3.26, "RESEARCH PROTOTYPE", fill=NAVY, color=WHITE, size=10)


def slide_one(prs: Presentation, team_id: str, team_name: str) -> None:
    slide = prs.slides[0]
    clear_slide_body(slide, title_slide=True)
    text(slide, 0.70, 1.20, 6.72, 0.70, "NISHAN-PQ", size=34, color=NAVY, bold=True)
    text(
        slide,
        0.72,
        1.88,
        6.70,
        0.78,
        "At-decryption session provenance for air-gapped documents.",
        size=17,
        color=TEAL_DARK,
        bold=True,
    )
    box(slide, 0.68, 2.48, 6.72, 3.68, fill=WHITE, line=GREY)
    fields = [
        ("Problem Statement ID", "SIH26237"),
        (
            "Problem Statement Title",
            "Cryptographic Attribution and Immutable Decryption Provenance for Multi-Recipient Encrypted Document Distribution",
        ),
        ("Theme", "Blockchain & Cybersecurity"),
        ("PS Category", "Software"),
        ("Team ID", team_id),
        ("Team Name (Registered on portal)", team_name),
    ]
    y = 2.73
    for index, (label, value) in enumerate(fields):
        height = 0.78 if index == 1 else 0.47
        text(slide, 0.95, y, 2.30, height, label.upper(), size=9.2, color=MUTED, bold=True)
        value_color = RED if value.startswith("[") else INK
        text(slide, 3.22, y - 0.01, 3.82, height, value, size=11.3 if index == 1 else 12.4, color=value_color, bold=index != 1)
        y += height
    draw_document_hero(slide)
    pill(slide, 8.52, 6.00, 3.50, "EVIDENCE UPDATE: 11 SEPTEMBER 2026", fill=PALE_AMBER, color="7A5100", size=8.6)


def slide_two(
    prs: Presentation,
    team_name: str,
    tardos_trials: dict,
    end_to_end_trials: dict,
    physical: dict,
) -> None:
    slide = prs.slides[1]
    trial_summaries = tardos_trials["summary"]
    codebooks = int(tardos_trials["parameters"]["trials"])
    attack_count = len(trial_summaries)
    trial_total = sum(int(item["trials"]) for item in trial_summaries.values())
    all_five_total = sum(
        int(item["all_colluders_recovered"]) for item in trial_summaries.values()
    )
    innocent_total = sum(
        int(item["total_innocent_accusations"]) for item in trial_summaries.values()
    )
    end_to_end_summary = end_to_end_trials["summary"]
    end_to_end_count = int(end_to_end_summary["exact_session_attributions"])
    end_to_end_count_total = int(end_to_end_trials["parameters"]["trials"])
    end_to_end_false = int(
        end_to_end_summary["false_accused_codebook_rows_total"]
    )
    end_to_end_mean = float(end_to_end_summary["expected_score_mean"])
    end_to_end_sd = float(end_to_end_summary["expected_score_population_sd"])
    clear_slide_body(slide)
    section_header(slide, "IDEA TITLE / PROPOSED SOLUTION", team_name, 2)
    rich_text(
        slide,
        2.30,
        0.82,
        10.35,
        0.38,
        [
            ("NISHAN-PQ  ", 18, NAVY, True),
            ("Session-bound evidence with explicit abstention", 13, TEAL_DARK, True),
        ],
    )
    text(
        slide,
        0.66,
        1.26,
        12.02,
        0.50,
        "Bind visual evidence + optional layout evidence + a recipient signature + a co-located quorum receipt to each released session.",
        size=11.8,
        color=INK,
        bold=True,
    )
    flow = [
        ("Encrypt once", "One AES-256-GCM ciphertext; no duplicate source files."),
        ("Authorize", "ML-KEM-768 wraps the same content key for each recipient."),
        ("Dual mark", "Tardos rendering row + HMAC tag in live PDF text layout."),
        ("Sign + commit", "Recipient ML-DSA receipt enters a 3-of-4 offline quorum."),
        ("Trace or abstain", "Check available channels; verify signed evidence."),
    ]
    x = 0.64
    for index, (title_value, body) in enumerate(flow, 1):
        small_card(slide, x, 1.82, 2.25, 1.18, index, title_value, body, accent=TEAL if index != 5 else AMBER)
        x += 2.48
        if index < len(flow):
            arrow(slide, x - 0.18, 2.29, w=0.17, h=0.30)

    latest = physical["profiles"]["affine_bias_translation"]["akshay3.jpeg"]
    box(slide, 0.65, 3.27, 5.92, 2.73, fill=PALE_TEAL, line=None)
    text(slide, 0.92, 3.50, 5.38, 0.34, "DIGITAL DEVELOPMENT", size=14, color=NAVY, bold=True)
    text(slide, 0.92, 3.91, 5.38, 0.45, f"{end_to_end_count}/{end_to_end_count_total} exact JPEG-Q55 sessions", size=19, color=TEAL_DARK, bold=True)
    bullet_list(slide, 0.92, 4.43, 5.36, 1.10, [
        f"All 1,000 rows scored; {end_to_end_false} extra rows accused",
        f"Mean {end_to_end_mean:,.3f}; population SD {end_to_end_sd:,.3f}",
        f"{codebooks} codebooks × {attack_count} strategies = {trial_total} code-level cases; not physical trials",
    ], size=10.2, gap=3)
    box(slide, 6.76, 3.27, 5.92, 2.73, fill=PALE_AMBER, line=None)
    text(slide, 7.03, 3.50, 5.38, 0.34, "REAL CAPTURES", size=14, color=NAVY, bold=True)
    text(slide, 7.03, 3.91, 5.38, 0.45, f"Historical {physical['historical_recovered']}/{physical['captures']}  ·  Experimental {physical['baseline_recovered']}/{physical['captures']} per profile", size=15.5, color="8A5D00", bold=True)
    bullet_list(slide, 7.03, 4.43, 5.36, 1.10, [
        f"Latest akshay3: {latest['expected_score']:,.3f} vs {latest['conditional_null']['threshold']:,.3f}",
        "Same akshay3 capture recovered; refinement adds no capture",
        "Four captures · one printed single-page public fixture",
        "Visual code-row research; not signed end-to-end physical attribution",
    ], size=10.2, gap=3, bullet_color=AMBER)
    text(slide, 0.84, 6.08, 11.62, 0.25, "LIMIT: small, non-independent fixtures; these results are not a deployed reliability or population-safety guarantee.", size=9.5, color=RED, bold=True, align=PP_ALIGN.CENTER)


def architecture_node(slide, x, y, w, title_value, body, *, fill=WHITE, accent=TEAL):
    box(slide, x, y, w, 1.08, fill=fill, line=accent, line_width=1.3)
    text(slide, x + 0.10, y + 0.10, w - 0.20, 0.30, title_value, size=11.2, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    text(slide, x + 0.10, y + 0.45, w - 0.20, 0.48, body, size=8.9, color=MUTED, align=PP_ALIGN.CENTER)


def slide_three(prs: Presentation, team_name: str) -> None:
    slide = prs.slides[2]
    clear_slide_body(slide)
    section_header(slide, "TECHNICAL APPROACH", team_name, 3)
    text(slide, 2.30, 0.85, 9.80, 0.30, "Implemented path, trusted-process assumption and demonstrator boundary", size=14, color=TEAL_DARK, bold=True)
    box(slide, 0.58, 1.30, 12.18, 3.20, fill=PALE, line=None)
    pill(slide, 0.82, 1.51, 2.34, "TRUSTED VIEWER PROCESS", fill=NAVY, color=WHITE, size=8.8)
    pill(slide, 9.22, 1.51, 3.02, "CO-LOCATED 3-OF-4 DEMONSTRATOR", fill=PALE_AMBER, color="7A5100", size=8.1)
    nodes = [
        (0.82, 2.12, 1.68, "Sender", "AES-256-GCM\none ciphertext"),
        (2.76, 2.12, 1.82, "Key envelopes", "ML-KEM-768\nper recipient"),
        (4.84, 2.12, 1.82, "Trusted viewer", "decrypts first; Tardos +\noptional layout tag"),
        (6.92, 2.12, 1.82, "Signed receipt", "ML-DSA-65\nenrollment-bound"),
        (9.00, 2.12, 1.82, "Quorum demo", "co-located keys/replicas\nhash-linked history"),
        (11.08, 2.12, 1.42, "Marked copy", "same visible\ncontent"),
    ]
    for index, (x, y, w, title_value, body) in enumerate(nodes):
        architecture_node(slide, x, y, w, title_value, body, fill=WHITE, accent=TEAL if index < 5 else AMBER)
        if index < len(nodes) - 1:
            arrow(slide, x + w + 0.07, y + 0.38, w=0.16, h=0.28)
    # forensic return path
    box(slide, 4.84, 3.52, 5.98, 0.62, fill=PALE_AMBER, line=None)
    text(slide, 4.99, 3.66, 5.68, 0.30, "LEAK → reference + visual/layout score → trace / ABSTAIN → receipt verification", size=9.7, color=NAVY, bold=True, align=PP_ALIGN.CENTER)

    box(slide, 0.62, 4.76, 5.96, 1.50, fill=WHITE, line=GREY)
    text(slide, 0.86, 4.98, 5.48, 0.30, "CURRENT BOUNDARIES", size=13.3, color=NAVY, bold=True)
    bullet_list(
        slide,
        0.86,
        5.36,
        5.45,
        0.72,
        [
            "Witness is not enforced in trace/release; plaintext exists before marking",
            "Initial audit/row allocation is not transaction-wide; reference retained",
        ],
        size=10.6,
        gap=3,
    )
    box(slide, 6.78, 4.76, 5.96, 1.50, fill=WHITE, line=GREY)
    text(slide, 7.02, 4.98, 5.48, 0.30, "PROPOSED HARDENING — NOT IMPLEMENTED", size=12.5, color=NAVY, bold=True)
    bullet_list(
        slide,
        7.02,
        5.36,
        5.45,
        0.72,
        [
            "Protected custody + external checkpoints + transaction-wide allocation",
            "Freeze release-bound assurance policy; separate keys and operators",
        ],
        size=10.6,
        gap=3,
    )


def metric_badge(slide, x, y, w, number, caption, *, fill=PALE_TEAL, accent=TEAL_DARK):
    box(slide, x, y, w, 0.86, fill=fill, line=None)
    text(slide, x + 0.06, y + 0.08, w - 0.12, 0.34, number, size=18, color=accent, bold=True, align=PP_ALIGN.CENTER)
    text(slide, x + 0.08, y + 0.48, w - 0.16, 0.25, caption, size=8.5, color=MUTED, bold=True, align=PP_ALIGN.CENTER)


def risk_row(slide, y, title_value, truth, gate, severity="HIGH"):
    box(slide, 6.06, y, 6.54, 0.94, fill=WHITE, line=GREY)
    fill = PALE_RED if severity == "HIGH" else PALE_AMBER
    color = RED if severity == "HIGH" else "9A6900"
    pill(slide, 6.22, y + 0.15, 0.72, severity, fill=fill, color=color, size=8.4)
    text(slide, 7.10, y + 0.10, 1.50, 0.28, title_value, size=10.4, color=NAVY, bold=True)
    text(slide, 8.58, y + 0.10, 3.78, 0.28, truth, size=9.0, color=MUTED)
    text(slide, 7.10, y + 0.49, 5.24, 0.29, "UNRESOLVED  ·  " + gate, size=8.9, color=TEAL_DARK, bold=True)


def slide_four(
    prs: Presentation,
    team_name: str,
    measured: dict,
    tardos_benchmark: dict,
    tardos_trials: dict,
    end_to_end_trials: dict,
    physical: dict,
) -> None:
    slide = prs.slides[3]
    clear_slide_body(slide)
    section_header(slide, "FEASIBILITY AND VIABILITY", team_name, 4)
    box(slide, 0.62, 1.24, 5.16, 4.96, fill=PALE_TEAL, line=None)
    text(slide, 0.90, 1.49, 4.62, 0.34, "COMPACT MEASURED OUTCOMES", size=14.2, color=NAVY, bold=True)
    bullet_list(
        slide,
        0.90,
        1.91,
        4.60,
        1.85,
        [
            f"Digital: {end_to_end_trials['summary']['exact_session_attributions']}/{end_to_end_trials['parameters']['trials']} exact JPEG-Q55 sessions; all {end_to_end_trials['parameters']['roster_rows_scored_per_trial']:,} rows scored",
            f"Physical: historical {physical['historical_recovered']}/{physical['captures']}; conditional {physical['baseline_recovered']}/{physical['captures']} per profile, same capture",
            "Searchable text retained; PDF bytes are not identical and signatures are not proven preserved",
            "Conflict/removal fixtures abstain; witness detects the recorded rollback fixture",
        ],
        size=11.1,
        gap=6,
    )
    psnr_values = [
        float(item["released_document"]["psnr_db"])
        for item in tardos_benchmark["carrier"]["metrics_by_user"]
    ]
    psnr_mean = sum(psnr_values) / len(psnr_values)
    metric_badge(
        slide,
        0.90,
        4.05,
        1.22,
        f"{psnr_mean:.2f}",
        "MEAN PSNR · 5 PDFs",
    )
    metric_badge(
        slide,
        2.24,
        4.05,
        1.78,
        f"{tardos_benchmark['attack_count']}",
        "RECORDED DIGITAL CASES",
        fill=PALE_AMBER,
        accent="9A6900",
    )
    trial_summaries = tardos_trials["summary"]
    trial_total = sum(int(item["trials"]) for item in trial_summaries.values())
    all_five_total = sum(
        int(item["all_colluders_recovered"]) for item in trial_summaries.values()
    )
    jpeg_violation_rate = 100 * float(
        tardos_benchmark["attacks"]["single_jpeg_q55"]["marking_condition"][
            "violation_rate"
        ]
    )
    metric_badge(
        slide,
        4.14,
        4.05,
        1.26,
        f"{all_five_total}/{trial_total}",
        "CODE TRIALS · ALL 5",
    )
    pill(slide, 0.92, 5.24, 4.46, "EXPERIMENTAL EVIDENCE · NOT A RELIABILITY RATE", fill=NAVY, color=WHITE, size=8.5)
    text(slide, 0.94, 5.72, 4.42, 0.27, "120 code-level cases are not 120 physical trials.", size=9.0, color=MUTED, align=PP_ALIGN.CENTER)

    text(slide, 6.06, 1.25, 6.54, 0.38, "FOUR MATERIAL RISK GROUPS", size=13.0, color=NAVY, bold=True)
    risk_row(slide, 1.76, "Physical", "One page / four captures", "held-out print/camera matrix", "HIGH")
    risk_row(slide, 2.80, "Raster / capacity", "Low-capacity bypass; raster policy open", "fail/restrict release policy", "HIGH")
    risk_row(slide, 3.84, "Rollback / row race", "Witness separate; allocation not atomic", "external checkpoint + atomic row", "HIGH")
    risk_row(slide, 4.88, "Framing / removal", "Authority can re-mark; both carriers removable", "custody + negative attacks", "HIGH")
    box(slide, 6.06, 5.94, 6.54, 0.30, fill=PALE_RED, line=None)
    text(slide, 6.12, 5.97, 6.42, 0.21, "Recipient session evidence is not proof of who leaked it.", size=9.2, color=RED, bold=True, align=PP_ALIGN.CENTER, margin=0)


def benefit_card(slide, x, y, w, title_value, body, *, number=None, accent=TEAL):
    box(slide, x, y, w, 1.16, fill=WHITE, line=GREY)
    if number:
        text(slide, x + 0.14, y + 0.12, 0.62, 0.38, number, size=17, color=accent, bold=True)
        title_x = x + 0.78
        title_w = w - 0.92
    else:
        title_x = x + 0.16
        title_w = w - 0.32
    text(slide, title_x, y + 0.13, title_w, 0.30, title_value, size=11.5, color=NAVY, bold=True)
    text(slide, x + 0.16, y + 0.53, w - 0.32, 0.48, body, size=9.2, color=MUTED)


def slide_five(prs: Presentation, team_name: str) -> None:
    slide = prs.slides[4]
    clear_slide_body(slide)
    section_header(slide, "IMPACT AND BENEFITS", team_name, 5)
    box(slide, 0.62, 1.22, 12.08, 1.20, fill=NAVY, line=None)
    text(slide, 0.92, 1.47, 2.25, 0.28, "INVESTIGATIVE ASSISTANCE", size=10, color=AMBER, bold=True)
    text(slide, 0.92, 1.78, 3.35, 0.36, "Offline session leads with PQ-signed evidence", size=12.4, color=WHITE, bold=True)
    text(slide, 4.55, 1.47, 1.60, 0.28, "CURRENT POSITION", size=10, color=AMBER, bold=True)
    text(slide, 4.55, 1.78, 7.58, 0.36, "Research proposal — not ready for operational accusations", size=13.0, color=WHITE, bold=True)

    benefits = [
        ("Session-level lead", "Narrows an investigation to signed release sessions; it does not identify the human actor."),
        ("Offline operation", "The measured demo does not depend on cloud APIs, tokens or a public chain."),
        ("PQC mechanism", "ML-KEM and ML-DSA protect key delivery and release receipts in the prototype."),
        ("Evidence discipline", "Conflict, removal and proof boundaries are recorded instead of hidden."),
    ]
    for index, (title_value, body) in enumerate(benefits):
        benefit_card(slide, 0.62 + index * 3.08, 2.70, 2.84, title_value, body, number=f"0{index+1}", accent=TEAL if index != 2 else AMBER)

    text(slide, 0.66, 4.12, 5.10, 0.30, "THREE GATES BEFORE OPERATIONAL USE", size=13.2, color=NAVY, bold=True)
    phases = [
        ("GATE 1", "Security boundaries + defensible claim language", "Release policy; accusation limits"),
        ("GATE 2", "Independent documents, keys and recipients", "Multi-page + signed/tagged PDFs + negative attacks"),
        ("GATE 3", "Held-out physical comparisons", "Matched visibility + security budgets"),
    ]
    for index, (phase, scope, metric) in enumerate(phases):
        x = 0.66 + index * 4.08
        fill = PALE_TEAL if index == 0 else PALE_AMBER if index == 1 else PALE
        box(slide, x, 4.54, 3.80, 1.52, fill=fill, line=None)
        pill(slide, x + 0.18, 4.71, 1.30, phase, fill=NAVY, color=WHITE, size=8.5)
        text(slide, x + 0.18, 5.15, 3.44, 0.30, scope, size=10.0, color=NAVY, bold=True)
        text(slide, x + 0.18, 5.51, 3.44, 0.34, "REQUIRES: " + metric, size=8.5, color=MUTED)
        if index < 2:
            arrow(slide, x + 3.86, 5.11, w=0.17, h=0.28, color=AMBER)


def reference_item(slide, y, number, title_value, link_label):
    circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(0.72), Inches(y), Inches(0.28), Inches(0.28))
    circle.fill.solid()
    circle.fill.fore_color.rgb = rgb(TEAL)
    circle.line.fill.background()
    text(slide, 0.72, y, 0.28, 0.28, str(number), size=8.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
    text(slide, 1.12, y - 0.02, 5.47, 0.23, title_value, size=9.6, color=NAVY, bold=True)
    text(slide, 1.12, y + 0.22, 5.47, 0.22, link_label, size=7.8, color=MUTED)


def slide_six(prs: Presentation, team_name: str, repository_url: str) -> None:
    slide = prs.slides[5]
    clear_slide_body(slide)
    section_header(slide, "RESEARCH AND REFERENCES", team_name, 6)
    text(slide, 0.68, 1.17, 5.92, 0.30, "PRIMARY SOURCES & FOUNDATIONAL WORK", size=13.3, color=NAVY, bold=True)
    references = [
        ("Official SIH26237 problem statement", "sih.gov.in/sih2026PS"),
        ("NIST FIPS 203 / 204 — ML-KEM / ML-DSA", "csrc.nist.gov/pubs/fips/203/final · /fips/204/final"),
        ("Tardos + T-Tracer + text-layout / dual marks", "ACM / IEEE sources; full URLs in notes"),
        ("Physical/screen-camera baselines", "StegaStamp · arXiv:2304.12682"),
        ("TRACE scope boundary", "arXiv:2607.08400 concerns LLM-agent trajectories"),
        ("Saved physical + digital evidence", "research/evidence/.../scores.json · artifacts/nishan/*.json"),
        ("Public repository", repository_url),
    ]
    y = 1.62
    for index, (title_value, link) in enumerate(references, 1):
        reference_item(slide, y, index, title_value, link)
        y += 0.64

    box(slide, 6.84, 1.19, 5.83, 4.92, fill=PALE, line=None)
    text(slide, 7.12, 1.49, 5.26, 0.33, "PRIOR-ART BOUNDARY", size=14.0, color=NAVY, bold=True)
    box(slide, 7.12, 1.98, 5.28, 1.45, fill=WHITE, line=GREY)
    text(slide, 7.34, 2.18, 4.84, 0.29, "KNOWN / PRESCRIBED", size=10.3, color=MUTED, bold=True)
    text(slide, 7.34, 2.61, 4.84, 0.51, "SIH26237 prescribes decryption marking, PQ signatures and DLT. Tardos, layout/dual marks and physical baselines pre-exist.", size=10.4, color=INK, bold=True)
    arrow(slide, 9.53, 3.54, w=0.36, h=0.29, color=AMBER)
    box(slide, 7.12, 3.95, 5.28, 1.63, fill=PALE_TEAL, line=None)
    text(slide, 7.34, 4.16, 4.84, 0.29, "HONEST CONTRIBUTION", size=10.3, color=TEAL_DARK, bold=True)
    text(slide, 7.34, 4.57, 4.84, 0.77, "Measured integration plus failure analysis: session-bound dual-domain evidence, explicit abstention, and documented physical limitations—not exclusive novelty or best-in-world performance.", size=9.9, color=NAVY, bold=True)
    pill(slide, 7.14, 5.72, 5.22, "Shorter code / stronger carrier: unvalidated hypotheses", fill=NAVY, color=WHITE, size=8.3)


def delete_last_slide(prs: Presentation) -> None:
    slide_id = prs.slides._sldIdLst[-1]
    prs.part.drop_rel(slide_id.rId)
    prs.slides._sldIdLst.remove(slide_id)


def build(
    template: Path,
    evidence_chart: Path,
    metrics_path: Path,
    tardos_benchmark_path: Path,
    tardos_trials_path: Path,
    end_to_end_trials_path: Path,
    output: Path,
    team_id: str,
    team_name: str,
    *,
    physical_path: Path | None = None,
    repository_url: str = "[INSERT PUBLIC REPOSITORY LINK BEFORE UPLOAD]",
) -> None:
    prs = Presentation(template)
    measured = json.loads(metrics_path.read_text())
    tardos_benchmark = json.loads(tardos_benchmark_path.read_text())
    tardos_trials = json.loads(tardos_trials_path.read_text())
    end_to_end_trials = json.loads(end_to_end_trials_path.read_text())
    if physical_path is None:
        physical_path = (
            Path(__file__).resolve().parents[2]
            / "research"
            / "evidence"
            / "nishan-bias-physical-2026-09-10"
            / "scores.json"
        )
    physical = load_physical(physical_path)
    if len(prs.slides) != 7:
        raise ValueError(f"expected 7-slide official template, got {len(prs.slides)}")
    delete_last_slide(prs)
    slide_one(prs, team_id, team_name)
    slide_two(
        prs,
        team_name,
        tardos_trials,
        end_to_end_trials,
        physical,
    )
    slide_three(prs, team_name)
    slide_four(
        prs,
        team_name,
        measured,
        tardos_benchmark,
        tardos_trials,
        end_to_end_trials,
        physical,
    )
    slide_five(prs, team_name)
    slide_six(prs, team_name, repository_url)
    psnr_values = [
        float(item["released_document"]["psnr_db"])
        for item in tardos_benchmark["carrier"]["metrics_by_user"]
    ]
    trial_summaries = tardos_trials["summary"]
    digital = {
        "jpeg_exact": int(end_to_end_trials["summary"]["exact_session_attributions"]),
        "jpeg_trials": int(end_to_end_trials["parameters"]["trials"]),
        "jpeg_false": int(end_to_end_trials["summary"]["false_accused_codebook_rows_total"]),
        "jpeg_mean": float(end_to_end_trials["summary"]["expected_score_mean"]),
        "jpeg_sd": float(end_to_end_trials["summary"]["expected_score_population_sd"]),
        "roster_rows": int(end_to_end_trials["parameters"]["roster_rows_scored_per_trial"]),
        "attack_cases": int(tardos_benchmark["attack_count"]),
        "code_cases": sum(int(item["trials"]) for item in trial_summaries.values()),
        "psnr_mean": sum(psnr_values) / len(psnr_values),
    }
    notes = speaker_notes(
        physical,
        team_id=team_id,
        team_name=team_name,
        repository_url=repository_url,
        digital=digital,
    )
    for slide, note in zip(prs.slides, notes, strict=True):
        slide.notes_slide.notes_text_frame.text = note
    prs.core_properties.title = "NISHAN-PQ — SIH26237 Idea Submission"
    prs.core_properties.subject = "Smart India Hackathon 2026"
    prs.core_properties.author = team_name
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    notes_path = output.with_name(f"{output.stem}_speaker_notes.md")
    notes_path.write_text(
        "# NISHAN-PQ speaker notes and audit record\n\n"
        + "\n\n".join(f"## Slide {index}\n\n{note}" for index, note in enumerate(notes, 1))
        + "\n"
    )
    print(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--template",
        type=Path,
        default=Path("resources/SIH2026-IDEA-Presentation-Format.pptx"),
    )
    parser.add_argument(
        "--evidence-chart",
        type=Path,
        default=Path("artifacts/nishan/attribution-scores.png"),
    )
    parser.add_argument(
        "--metrics",
        type=Path,
        default=Path("artifacts/nishan/measured-results.json"),
    )
    parser.add_argument(
        "--tardos-benchmark",
        type=Path,
        default=Path("artifacts/nishan/tardos-pdf-benchmark.json"),
    )
    parser.add_argument(
        "--tardos-trials",
        type=Path,
        default=Path("artifacts/nishan/tardos-independent-codebook-study.json"),
    )
    parser.add_argument(
        "--end-to-end-trials",
        type=Path,
        default=Path("artifacts/nishan/end-to-end-jpeg-trials.json"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("submissions/SIH26237_NISHAN_PQ/NISHAN-PQ_SIH26237.pptx"),
    )
    parser.add_argument("--team-id", default="[ENTER TEAM ID]")
    parser.add_argument("--team-name", default="[ENTER REGISTERED TEAM NAME]")
    parser.add_argument(
        "--physical-evidence",
        type=Path,
        default=None,
        help="validated physical score JSON; defaults to the fixed documented evidence",
    )
    parser.add_argument(
        "--repository-url",
        default="[INSERT PUBLIC REPOSITORY LINK BEFORE UPLOAD]",
    )
    args = parser.parse_args()
    build(
        args.template.resolve(),
        args.evidence_chart.resolve(),
        args.metrics.resolve(),
        args.tardos_benchmark.resolve(),
        args.tardos_trials.resolve(),
        args.end_to_end_trials.resolve(),
        args.output.resolve(),
        args.team_id,
        args.team_name,
        physical_path=args.physical_evidence.resolve() if args.physical_evidence else None,
        repository_url=args.repository_url,
    )


if __name__ == "__main__":
    main()
