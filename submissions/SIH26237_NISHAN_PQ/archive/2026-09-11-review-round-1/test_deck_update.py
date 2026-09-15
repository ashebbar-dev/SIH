from __future__ import annotations

import copy
import importlib
import json
import math
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

import pymupdf as fitz
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[3]
PACKAGE = ROOT / "submissions" / "SIH26237_NISHAN_PQ"
PHYSICAL = ROOT / "research" / "evidence" / "nishan-bias-physical-2026-09-10" / "scores.json"
sys.path.insert(0, str(PACKAGE))


def slide_text(prs: Presentation) -> list[str]:
    return [
        "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text_frame"))
        for slide in prs.slides
    ]


def note_text(slide) -> str:
    return slide.notes_slide.notes_text_frame.text.strip()


class PhysicalEvidenceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.deck_evidence = importlib.import_module("deck_evidence")
        self.source = json.loads(PHYSICAL.read_text())

    def write_mutation(self, mutator) -> Path:
        data = copy.deepcopy(self.source)
        mutator(data)
        handle = tempfile.NamedTemporaryFile(mode="w", suffix=".json", delete=False)
        with handle:
            json.dump(data, handle, allow_nan=True)
        self.addCleanup(Path(handle.name).unlink, missing_ok=True)
        return Path(handle.name)

    def test_exact_physical_summary(self) -> None:
        summary = self.deck_evidence.load_physical(PHYSICAL)
        self.assertEqual(summary["captures"], 4)
        self.assertEqual(summary["historical_recovered"], 0)
        self.assertEqual(summary["baseline_recovered"], 1)
        self.assertEqual(summary["refined_recovered"], 1)
        self.assertEqual(summary["boundary_failures"], 2)
        self.assertEqual(set(summary["profiles"]), {"affine_baseline", "affine_bias_translation"})
        self.assertEqual(len(summary["records"]), 8)

    def test_missing_input_is_explicit(self) -> None:
        with self.assertRaises(FileNotFoundError):
            self.deck_evidence.load_physical(PHYSICAL.with_name("missing.json"))

    def test_duplicate_record_is_rejected(self) -> None:
        path = self.write_mutation(lambda data: data["records"].append(copy.deepcopy(data["records"][0])))
        with self.assertRaisesRegex(ValueError, "duplicate"):
            self.deck_evidence.load_physical(path)

    def test_missing_profile_is_rejected(self) -> None:
        def mutate(data):
            data["records"] = [r for r in data["records"] if r["profile"] != "affine_bias_translation"]

        with self.assertRaisesRegex(ValueError, "profile|records"):
            self.deck_evidence.load_physical(self.write_mutation(mutate))

    def test_scored_record_requires_finite_threshold(self) -> None:
        for bad_value in (None, math.nan):
            with self.subTest(value=bad_value):
                def mutate(data, value=bad_value):
                    data["records"][0]["conditional_null"]["threshold"] = value

                with self.assertRaisesRegex(ValueError, "threshold"):
                    self.deck_evidence.load_physical(self.write_mutation(mutate))

    def test_false_counted_success_is_rejected(self) -> None:
        def mutate(data):
            record = data["records"][0]
            record["conditional_accused"] = [record["expected_row"]]
            record["row0_accused"] = True

        with self.assertRaisesRegex(ValueError, "accus|threshold|success"):
            self.deck_evidence.load_physical(self.write_mutation(mutate))


class DeckAndVerifierTests(unittest.TestCase):
    maxDiff = None

    @classmethod
    def setUpClass(cls) -> None:
        cls.build_deck = importlib.import_module("build_deck")
        cls.verify_submission = importlib.import_module("verify_submission")

    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.output_dir = self.root / "submissions" / "SIH26237_NISHAN_PQ"
        self.output_dir.mkdir(parents=True)
        self.pptx = self.output_dir / "NISHAN-PQ_SIH26237.pptx"
        self.pdf = self.output_dir / "NISHAN-PQ_SIH26237.pdf"
        self.notes = self.output_dir / "NISHAN-PQ_SIH26237_speaker_notes.md"
        self.team_id = "TEAM-26237"
        self.team_name = "NISHAN TEST TEAM"
        self.repository_url = "https://example.org/nishan/source"

        self.build_deck.build(
            ROOT / "resources" / "SIH2026-IDEA-Presentation-Format.pptx",
            ROOT / "artifacts" / "nishan" / "attribution-scores.png",
            ROOT / "artifacts" / "nishan" / "measured-results.json",
            ROOT / "artifacts" / "nishan" / "tardos-pdf-benchmark.json",
            ROOT / "artifacts" / "nishan" / "tardos-independent-codebook-study.json",
            ROOT / "artifacts" / "nishan" / "end-to-end-jpeg-trials.json",
            self.pptx,
            self.team_id,
            self.team_name,
            physical_path=PHYSICAL,
            repository_url=self.repository_url,
        )

    def prepare_verifier_root(self) -> None:
        artifact_dir = self.root / "artifacts" / "nishan"
        artifact_dir.mkdir(parents=True)
        for name in (
            "measured-results.json",
            "tardos-pdf-benchmark.json",
            "tardos-independent-codebook-study.json",
            "end-to-end-jpeg-trials.json",
        ):
            shutil.copy2(ROOT / "artifacts" / "nishan" / name, artifact_dir / name)
        physical_dir = self.root / "research" / "evidence" / "nishan-bias-physical-2026-09-10"
        physical_dir.mkdir(parents=True)
        shutil.copy2(PHYSICAL, physical_dir / "scores.json")
        self.render_pdf()

    def render_pdf(self, stale: str = "") -> None:
        prs = Presentation(self.pptx)
        document = fitz.open()
        for index, content in enumerate(slide_text(prs)):
            page = document.new_page(width=960, height=540)
            prefix = stale + "\n" if index == 0 and stale else ""
            page.insert_textbox(fitz.Rect(18, 18, 942, 522), prefix + content, fontsize=5.8)
        document.save(self.pdf)
        document.close()

    def test_build_has_six_editable_slides_and_complete_matching_notes(self) -> None:
        prs = Presentation(self.pptx)
        self.assertEqual(len(prs.slides), 6)
        texts = slide_text(prs)
        combined = "\n".join(texts)
        self.assertIn(self.team_id, combined)
        self.assertIn(self.team_name, combined)
        self.assertIn(self.repository_url, combined)
        self.assertNotRegex(combined, r"\[(?:ENTER|INSERT)[^]]+\]")
        self.assertNotIn("real physical captures pending", combined)
        self.assertNotIn("Synthetic perspective passes; hardware pending", combined)
        embedded = [note_text(slide) for slide in prs.slides]
        self.assertTrue(all(embedded))
        exported = self.notes.read_text()
        for index, block in enumerate(embedded, 1):
            self.assertIn(f"## Slide {index}", exported)
            self.assertIn(block, exported)
        joined = "\n".join(embedded)
        for prefix, count in (("R", 14), ("P", 24)):
            for number in range(1, count + 1):
                self.assertEqual(joined.count(f"{prefix}{number}:"), 1)
        semantic_anchors = {
            "R8:": "NOT MET",
            "R14:": "No public blockchain",
            "P1:": "separate witness",
            "P3:": "masquerading as a raster/image",
            "P10:": "same akshay3 capture recovered in both—not two recovered captures",
            "P11:": "one person across 30 accumulated issued sessions",
            "P23:": "stray file '=4.11'",
            "P24:": "Novelty is unestablished",
        }
        for label, anchor in semantic_anchors.items():
            line = next(item for item in joined.splitlines() if item.startswith(label))
            self.assertIn(anchor.lower(), line.lower())

    def test_slide_claims_are_evidence_derived(self) -> None:
        combined = "\n".join(slide_text(Presentation(self.pptx)))
        self.assertIn("30/30", combined)
        self.assertIn("0/4", combined)
        self.assertIn("1/4", combined)
        self.assertIn("Experimental 1/4 per profile", combined)
        self.assertIn("Same akshay3 capture recovered; refinement adds no capture", combined)
        self.assertIn("Check available channels; verify signed evidence", combined)
        self.assertIn("Multi-page + signed/tagged PDFs + negative attacks", combined)
        self.assertIn("1,827.334", combined)
        self.assertIn("1,096.290", combined)
        self.assertIn("Recipient session evidence is not proof of who leaked it.", combined)

    def test_verifier_rejects_stale_pptx_claim(self) -> None:
        self.prepare_verifier_root()
        prs = Presentation(self.pptx)
        prs.slides[0].shapes.add_textbox(Inches(0.2), Inches(0.2), Inches(5), Inches(0.3)).text = (
            "real physical captures pending"
        )
        prs.save(self.pptx)
        result = self.verify_submission.verify(self.root, allow_placeholders=True)
        self.assertFalse(result["ok"])
        self.assertTrue(any("stale" in error.lower() and "PPTX" in error for error in result["errors"]))

    def test_verifier_rejects_stale_pdf_claim(self) -> None:
        self.prepare_verifier_root()
        self.render_pdf("Synthetic perspective passes; hardware pending")
        result = self.verify_submission.verify(self.root, allow_placeholders=True)
        self.assertFalse(result["ok"])
        self.assertTrue(any("stale" in error.lower() and "PDF" in error for error in result["errors"]))

    def test_verifier_rejects_missing_physical_evidence(self) -> None:
        self.prepare_verifier_root()
        physical = self.root / "research" / "evidence" / "nishan-bias-physical-2026-09-10" / "scores.json"
        physical.unlink()
        result = self.verify_submission.verify(self.root, allow_placeholders=True)
        self.assertFalse(result["ok"])
        self.assertTrue(any("physical evidence" in error.lower() and "missing" in error.lower() for error in result["errors"]))


if __name__ == "__main__":
    unittest.main()
