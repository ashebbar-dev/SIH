#!/usr/bin/env python3
"""Verify the structural and evidence claims of the NISHAN SIH package."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import pymupdf as fitz
from pptx import Presentation


REQUIRED_TEXT = [
    "SIH26237",
    "Cryptographic Attribution",
    "PROPOSED SOLUTION",
    "TECHNICAL APPROACH",
    "FEASIBILITY AND VIABILITY",
    "IMPACT AND BENEFITS",
    "RESEARCH AND REFERENCES",
    "ML-KEM-768",
    "ML-DSA-65",
    "enrollment",
    "52,500",
    "1,000",
    "HMAC",
    "transplant",
    "30 codebooks",
    "30/30",
    "JPEG-Q55 sessions",
    "CODE TRIALS",
    "Both carriers remain removable",
    "hardware pending",
]


def verify(root: Path, allow_placeholders: bool) -> dict[str, object]:
    pdf = root / "submissions" / "SIH26237_NISHAN_PQ" / "NISHAN-PQ_SIH26237.pdf"
    pptx = root / "submissions" / "SIH26237_NISHAN_PQ" / "NISHAN-PQ_SIH26237.pptx"
    evidence_path = root / "artifacts" / "nishan" / "measured-results.json"
    tardos_benchmark_path = root / "artifacts" / "nishan" / "tardos-pdf-benchmark.json"
    tardos_trials_path = root / "artifacts" / "nishan" / "tardos-independent-codebook-study.json"
    end_to_end_trials_path = root / "artifacts" / "nishan" / "end-to-end-jpeg-trials.json"
    errors: list[str] = []
    warnings: list[str] = []

    if not pdf.exists():
        errors.append(f"missing PDF: {pdf}")
        pdf_text = ""
        normalized_pdf_text = ""
        pdf_pages = 0
    else:
        document = fitz.open(pdf)
        pdf_pages = len(document)
        pdf_text = "\n".join(page.get_text() for page in document)
        normalized_pdf_text = re.sub(r"\s+", " ", pdf_text)
        page_sizes = {(round(page.rect.width, 1), round(page.rect.height, 1)) for page in document}
        document.close()
        if pdf_pages != 6:
            errors.append(f"PDF has {pdf_pages} pages; SIH requires six")
        if page_sizes != {(960.0, 540.0)}:
            errors.append(f"unexpected PDF page sizes: {sorted(page_sizes)}")
        for phrase in REQUIRED_TEXT:
            if phrase not in normalized_pdf_text:
                errors.append(f"PDF is missing required text: {phrase!r}")

    if not pptx.exists():
        errors.append(f"missing editable deck: {pptx}")
        pptx_slides = 0
    else:
        presentation = Presentation(pptx)
        pptx_slides = len(presentation.slides)
        if pptx_slides != 6:
            errors.append(f"PPTX has {pptx_slides} slides; expected six")

    placeholders = sorted(set(re.findall(r"\[[A-Z][A-Z _-]+\]", pdf_text)))
    if placeholders:
        message = "replace before upload: " + ", ".join(placeholders)
        if allow_placeholders:
            warnings.append(message)
        else:
            errors.append(message)

    evidence = json.loads(evidence_path.read_text())
    jpeg = evidence.get("jpeg_attack_attribution", [])
    if len(jpeg) != 1 or jpeg[0].get("recipient_id") != "bob":
        errors.append("measured JPEG evidence is not a unique Bob attribution")
    elif not jpeg[0].get("recipient_signature_valid"):
        errors.append("Bob attribution has no valid recipient signature")
    coalition = evidence.get("two_copy_average_attribution", [])
    coalition_ids = {item.get("recipient_id") for item in coalition}
    if coalition_ids != {"alice", "bob"}:
        errors.append(f"measured coalition attribution is {sorted(coalition_ids)}")
    if not all(item.get("recipient_signature_valid") for item in coalition):
        errors.append("one or more coalition attributions has an invalid signature")
    three_way = evidence.get("three_copy_average_attribution", [])
    three_way_ids = {item.get("recipient_id") for item in three_way}
    if three_way_ids != {"alice", "bob", "charlie"}:
        errors.append(f"measured three-copy attribution is {sorted(three_way_ids)}")
    if not all(item.get("recipient_signature_valid") for item in three_way):
        errors.append("one or more three-copy attributions has an invalid signature")
    calibrations = evidence.get("threshold_calibration", {})
    for attack in ("jpeg_q55", "two_copy_average", "three_copy_average"):
        calibration = calibrations.get(attack, {})
        if int(calibration.get("roster_size", 0)) != 1_000:
            errors.append(f"{attack} does not score the declared 1,000-row roster")
        if int(calibration.get("code_length", 0)) != 52_500:
            errors.append(f"{attack} does not use the 52,500-symbol profile")
        if float(calibration.get("threshold", 0)) != 2_100:
            errors.append(f"{attack} does not use Tardos threshold 2,100")
        if calibration.get("formal_code_bound_available") is not True:
            errors.append(f"{attack} omits the reference code-bound metadata")
        if calibration.get("formal_familywise_bound_established") is not False:
            errors.append(f"{attack} overstates physical-attack theorem applicability")
    ledger = evidence.get("ledger_after_single_replica_attack", {})
    if ledger.get("quorum_valid") is not True or ledger.get("divergent_replicas") != ["validator-1"]:
        errors.append("single-replica ledger attack evidence does not retain and identify quorum")
    invalid_receipt = evidence.get("invalid_recipient_signature_attack", {})
    if invalid_receipt.get("invalid_record_received_quorum_endorsement") is not True:
        errors.append("invalid-receipt fixture did not reach the application validation boundary")
    if invalid_receipt.get("trace_blocked") is not True:
        errors.append("tracing did not fail closed on an invalid recipient receipt")
    if invalid_receipt.get("future_plaintext_release_blocked") is not True:
        errors.append("future plaintext release did not fail closed on an invalid recipient receipt")
    if invalid_receipt.get("plaintext_output_created") is not False:
        errors.append("plaintext was created after an invalid recipient receipt")
    substituted_identity = evidence.get("self_asserted_identity_key_attack", {})
    if substituted_identity.get("substituted_signature_valid_under_attacker_key") is not True:
        errors.append("identity-substitution fixture did not contain a valid attacker-key signature")
    if substituted_identity.get("invalid_record_received_quorum_endorsement") is not True:
        errors.append("identity-substitution fixture did not reach the enrollment check")
    if substituted_identity.get("trace_blocked") is not True:
        errors.append("tracing accepted a self-asserted key for an enrolled identity")
    if substituted_identity.get("future_plaintext_release_blocked") is not True:
        errors.append("release accepted a self-asserted key for an enrolled identity")
    if substituted_identity.get("plaintext_output_created") is not False:
        errors.append("plaintext was created after identity-key substitution")
    registry_tamper = evidence.get("enrollment_registry_tamper_attack", {})
    if registry_tamper.get("trust_root_valid") is not False:
        errors.append("edited enrollment registry retained a valid trust root")
    if registry_tamper.get("quorum_valid") is not False:
        errors.append("edited enrollment registry retained a valid ledger quorum")
    if registry_tamper.get("future_plaintext_release_blocked") is not True:
        errors.append("release continued after enrollment-registry tampering")
    if registry_tamper.get("plaintext_output_created") is not False:
        errors.append("plaintext was created after enrollment-registry tampering")
    enrollment_root = evidence.get("enrollment_trust_root", {})
    if enrollment_root.get("trust_root_valid") is not True:
        errors.append("baseline enrollment trust root is not valid")
    if enrollment_root.get("enrolled_identity_ids") != ["alice", "bob", "charlie"]:
        errors.append("baseline enrollment registry does not contain the three demo identities")
    replay = evidence.get("valid_receipt_replay_attack", {})
    if replay.get("replayed_record_received_quorum_endorsement") is not True:
        errors.append("valid-receipt replay fixture did not reach the application gate")
    if replay.get("trace_blocked") is not True:
        errors.append("tracing accepted a replayed valid receipt")
    if replay.get("future_plaintext_release_blocked") is not True:
        errors.append("release accepted a replayed valid receipt")
    if replay.get("plaintext_output_created") is not False:
        errors.append("plaintext was created after a valid-receipt replay")
    rollback = evidence.get("coordinated_valid_prefix_rollback_attack", {})
    if rollback.get("baseline_witness_comparison") != "consistent":
        errors.append("external witness does not match the baseline ledger head")
    if rollback.get("rolled_back_ledger_quorum_valid") is not True:
        errors.append("rollback fixture did not preserve an apparently valid ledger quorum")
    if rollback.get("external_witness_comparison") != "rollback_detected":
        errors.append("external witness failed to detect coordinated valid-prefix rollback")
    if not int(rollback.get("rolled_back_length", 0)) < int(
        rollback.get("witnessed_length", 0)
    ):
        errors.append("rollback fixture did not shorten the witnessed ledger")

    # The submission-safe artifact directory must not contain keys or watermark secrets.
    suspicious: list[str] = []
    for path in (root / "artifacts" / "nishan").rglob("*"):
        if not path.is_file():
            continue
        lowered = path.name.lower()
        if "private" in lowered or "secret" in lowered or path.suffix.lower() == ".pem":
            suspicious.append(str(path.relative_to(root)))
            continue
        if path.stat().st_size < 5_000_000:
            raw = path.read_bytes()
            if b"BEGIN PRIVATE KEY" in raw or b"watermark-secret" in raw:
                suspicious.append(str(path.relative_to(root)))
    if suspicious:
        errors.append("submission-safe artifacts contain secret-like files: " + ", ".join(suspicious))

    benchmark = json.loads(tardos_benchmark_path.read_text())
    five_copy = benchmark["attacks"]["5_copy_pixel_average"]
    if five_copy.get("colluders_recovered_count") != 5:
        errors.append("five-copy carrier fixture did not recover all five coalition rows")
    if five_copy.get("innocent_accusations_count") != 0:
        errors.append("five-copy carrier fixture accused an innocent roster row")
    if five_copy.get("marking_condition", {}).get("violations") != 0:
        errors.append("five-copy fixture does not satisfy the recorded marking condition")
    jpeg_benchmark = benchmark["attacks"]["single_jpeg_q55"]
    if jpeg_benchmark.get("colluders_recovered_count") != 1:
        errors.append("full carrier benchmark did not recover the JPEG source row")
    if jpeg_benchmark.get("marking_condition", {}).get("violations", 0) <= 0:
        errors.append("JPEG benchmark hides its marking-condition violations")
    removed = benchmark["attacks"]["pdf_overlay_object_removed"]
    if removed.get("colluders_recovered_count") != 0:
        errors.append("overlay-removal red-team fixture unexpectedly attributes a row")
    if removed.get("fusion", {}).get("decision") != "abstain_single_channel_editable_pdf":
        errors.append("overlay-removal fixture does not fail closed")
    direct_pdf = benchmark["attacks"]["single_digital_pdf"]
    if direct_pdf.get("fusion", {}).get("decision") != "corroborated_channels":
        errors.append("direct PDF does not produce dual-channel corroboration")
    transplant = benchmark["attacks"]["pdf_visual_carrier_transplanted"]
    if transplant.get("fusion", {}).get("decision") != "abstain_channel_conflict":
        errors.append("cross-recipient carrier transplant does not force abstention")
    transplant_stripped = benchmark["attacks"]["pdf_visual_transplant_then_layout_strip"]
    if (
        transplant_stripped.get("fusion", {}).get("decision")
        != "abstain_single_channel_editable_pdf"
    ):
        errors.append("transplant-then-strip attack does not force abstention")
    partial_overlap = benchmark["attacks"][
        "pdf_partial_overlap_visual_coalition_on_bob_layout"
    ]
    if partial_overlap.get("fusion", {}).get("decision") != "abstain_channel_conflict":
        errors.append("partial-overlap carrier recombination does not force abstention")
    if partial_overlap.get("fusion", {}).get("selected_indices") != []:
        errors.append("partial-overlap carrier recombination selected an identity")
    both_removed = benchmark["attacks"]["pdf_both_channels_removed"]
    if both_removed.get("fusion", {}).get("decision") != "no_attribution_signal":
        errors.append("both-channel removal fixture does not expose signal loss")
    perspective = benchmark["attacks"]["single_synthetic_perspective_photo"]
    if perspective.get("colluders_recovered_count") != 1:
        errors.append("synthetic perspective fixture did not recover its source row")
    if not perspective.get("decoder_preprocessing", {}).get("pages", [{}])[0].get("applied"):
        errors.append("synthetic perspective fixture did not exercise registration")
    if benchmark.get("attack_count") != 25 or benchmark.get("all_attack_expectations_passed") is not True:
        errors.append("the declared 25-case attack matrix is incomplete or failed")
    if benchmark["carrier"]["live_document_checks"].get("text_extraction_equal") is not True:
        errors.append("marked carrier benchmark does not retain PDF text")

    trials_study = json.loads(tardos_trials_path.read_text())
    if int(trials_study.get("parameters", {}).get("trials", 0)) != 30:
        errors.append("independent-codebook study does not contain 30 trials")
    code_trial_total = 0
    all_five_total = 0
    for attack in ("interleaving", "majority", "minority", "coin_flip"):
        summary = trials_study.get("summary", {}).get(attack, {})
        attack_trials = int(summary.get("trials", 0))
        code_trial_total += attack_trials
        all_five_total += int(summary.get("all_colluders_recovered", 0))
        if int(summary.get("total_innocent_accusations", -1)) != 0:
            errors.append(f"{attack} codebook study contains an innocent accusation")
        if int(summary.get("all_colluders_recovered", 0)) != attack_trials:
            errors.append(f"{attack} did not recover all five colluders in every recorded trial")

    end_to_end_study = json.loads(end_to_end_trials_path.read_text())
    end_to_end_parameters = end_to_end_study.get("parameters", {})
    end_to_end_summary = end_to_end_study.get("summary", {})
    if int(end_to_end_parameters.get("trials", 0)) != 30:
        errors.append("repeated end-to-end JPEG study does not contain 30 trials")
    if int(end_to_end_parameters.get("roster_rows_scored_per_trial", 0)) != 1_000:
        errors.append("repeated end-to-end JPEG study does not score all 1,000 rows")
    if int(end_to_end_summary.get("exact_session_attributions", 0)) != 30:
        errors.append("repeated end-to-end JPEG study missed an exact session")
    if int(end_to_end_summary.get("valid_recipient_signatures", 0)) != 30:
        errors.append("repeated end-to-end JPEG study contains an invalid receipt")
    if int(end_to_end_summary.get("false_accused_codebook_rows_total", -1)) != 0:
        errors.append("repeated end-to-end JPEG study accused an extra codebook row")
    if float(end_to_end_summary.get("expected_score_min", 0)) <= float(
        end_to_end_parameters.get("tardos_threshold", 2_100)
    ):
        errors.append("a repeated end-to-end JPEG score did not cross threshold")

    expected_score = f"{float(jpeg[0]['score']):,.0f}" if jpeg else ""

    result = {
        "ok": not errors,
        "pdf_pages": pdf_pages,
        "pptx_slides": pptx_slides,
        "measured_jpeg_score": float(jpeg[0]["score"]) if jpeg else None,
        "measured_jpeg_score_display": expected_score,
        "tardos_profile": {"n": 1000, "c": 5, "m": 52500, "Z": 2100},
        "five_copy_recovered": five_copy.get("colluders_recovered_count"),
        "overlay_removal_recovered": removed.get("colluders_recovered_count"),
        "dual_carrier_direct_decision": direct_pdf.get("fusion", {}).get("decision"),
        "transplant_decision": transplant.get("fusion", {}).get("decision"),
        "attack_matrix_cases": benchmark.get("attack_count"),
        "code_level_trial_pairs": code_trial_total,
        "code_level_all_five_recovered": all_five_total,
        "end_to_end_exact_sessions": end_to_end_summary.get(
            "exact_session_attributions"
        ),
        "end_to_end_score_mean": end_to_end_summary.get("expected_score_mean"),
        "end_to_end_score_population_sd": end_to_end_summary.get(
            "expected_score_population_sd"
        ),
        "coalition_recipients": sorted(coalition_ids),
        "three_copy_recipients": sorted(three_way_ids),
        "placeholders": placeholders,
        "warnings": warnings,
        "errors": errors,
    }
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[2])
    parser.add_argument(
        "--allow-placeholders",
        action="store_true",
        help="report team/repository placeholders as warnings while details are unavailable",
    )
    args = parser.parse_args()
    result = verify(args.root.resolve(), args.allow_placeholders)
    print(json.dumps(result, indent=2))
    raise SystemExit(0 if result["ok"] else 1)


if __name__ == "__main__":
    main()
