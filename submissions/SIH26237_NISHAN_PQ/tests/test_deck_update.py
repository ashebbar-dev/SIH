from __future__ import annotations

import copy
import importlib
import json
import math
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

import pymupdf as fitz
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[3]
PACKAGE = ROOT / "submissions" / "SIH26237_NISHAN_PQ"
PHYSICAL = ROOT / "research" / "evidence" / "nishan-bias-physical-2026-09-10" / "scores.json"
SAFEGUARDS = ROOT / "research" / "evidence" / "nishan-selection-2026-09-12" / "demo" / "results.json"
sys.path.insert(0, str(PACKAGE))


def slide_text(prs: Presentation) -> list[str]:
    return [
        "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text_frame"))
        for slide in prs.slides
    ]


def note_text(slide) -> str:
    return slide.notes_slide.notes_text_frame.text.strip()


class PhysicalEvidenceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.deck_evidence = importlib.import_module("deck_evidence")
        self.source = json.loads(PHYSICAL.read_text())

    def write_mutation(self, mutator) -> Path:
        data = copy.deepcopy(self.source)
        mutator(data)
        handle = tempfile.NamedTemporaryFile(mode="w", suffix=".json", delete=False)
        with handle:
            json.dump(data, handle, allow_nan=True)
        self.addCleanup(Path(handle.name).unlink, missing_ok=True)
        return Path(handle.name)

    def test_exact_physical_summary(self) -> None:
        summary = self.deck_evidence.load_physical(PHYSICAL)
        self.assertEqual(summary["captures"], 4)
        self.assertEqual(summary["historical_recovered"], 0)
        self.assertEqual(summary["baseline_recovered"], 1)
        self.assertEqual(summary["refined_recovered"], 1)
        self.assertEqual(summary["boundary_failures"], 2)
        self.assertEqual(set(summary["profiles"]), {"affine_baseline", "affine_bias_translation"})
        self.assertEqual(len(summary["records"]), 8)

    def test_missing_input_is_explicit(self) -> None:
        with self.assertRaises(FileNotFoundError):
            self.deck_evidence.load_physical(PHYSICAL.with_name("missing.json"))

    def test_duplicate_record_is_rejected(self) -> None:
        path = self.write_mutation(lambda data: data["records"].append(copy.deepcopy(data["records"][0])))
        with self.assertRaisesRegex(ValueError, "duplicate"):
            self.deck_evidence.load_physical(path)

    def test_missing_profile_is_rejected(self) -> None:
        def mutate(data):
            data["records"] = [r for r in data["records"] if r["profile"] != "affine_bias_translation"]

        with self.assertRaisesRegex(ValueError, "profile|records"):
            self.deck_evidence.load_physical(self.write_mutation(mutate))

    def test_scored_record_requires_finite_threshold(self) -> None:
        for bad_value in (None, math.nan):
            with self.subTest(value=bad_value):
                def mutate(data, value=bad_value):
                    data["records"][0]["conditional_null"]["threshold"] = value

                with self.assertRaisesRegex(ValueError, "threshold"):
                    self.deck_evidence.load_physical(self.write_mutation(mutate))

    def test_false_counted_success_is_rejected(self) -> None:
        def mutate(data):
            record = data["records"][0]
            record["conditional_accused"] = [record["expected_row"]]
            record["row0_accused"] = True

        with self.assertRaisesRegex(ValueError, "accus|threshold|success"):
            self.deck_evidence.load_physical(self.write_mutation(mutate))

    def test_unaccused_competing_score_above_historical_threshold_is_rejected(self) -> None:
        def mutate(data):
            record = next(
                item
                for item in data["records"]
                if item["capture"] == "akshay.jpeg" and item["profile"] == "affine_baseline"
            )
            record["highest_other_score"] = 2_200.0
            record["conditional_null"]["threshold"] = 3_000.0

        with self.assertRaisesRegex(ValueError, "historical.*competing|competing.*historical"):
            self.deck_evidence.load_physical(self.write_mutation(mutate))

    def test_renamed_capture_is_rejected_explicitly(self) -> None:
        def mutate(data):
            data["settings"]["captures"][-1] = "renamed.jpeg"
            for record in data["records"]:
                if record["capture"] == "akshay3.jpeg":
                    record["capture"] = "renamed.jpeg"

        with self.assertRaisesRegex(ValueError, "capture names.*akshay3"):
            self.deck_evidence.load_physical(self.write_mutation(mutate))

    def test_malformed_identity_names_raise_explicit_value_error(self) -> None:
        mutations = {
            "capture settings": lambda data: data["settings"]["captures"].__setitem__(0, {}),
            "profile settings": lambda data: data["settings"]["profiles"].__setitem__(0, []),
            "record identity": lambda data: data["records"][0].__setitem__("capture", {}),
        }
        for label, mutate in mutations.items():
            with self.subTest(label=label):
                with self.assertRaisesRegex(ValueError, "name|identity|capture|profile"):
                    self.deck_evidence.load_physical(self.write_mutation(mutate))

    def test_recovery_swapped_to_wrong_capture_is_rejected_explicitly(self) -> None:
        def mutate(data):
            records = {
                (record["capture"], record["profile"]): record
                for record in data["records"]
            }
            first = records[("akshay2.jpeg", "affine_baseline")]
            second = records[("akshay3.jpeg", "affine_baseline")]
            outcome_fields = (
                "status",
                "expected_row",
                "expected_score",
                "highest_other_score",
                "bit_error_fraction",
                "historical_accused",
                "conditional_accused",
                "conditional_null",
                "row0_accused",
            )
            for field in outcome_fields:
                first[field], second[field] = second[field], first[field]

        with self.assertRaisesRegex(ValueError, "recovery identity.*akshay3"):
            self.deck_evidence.load_physical(self.write_mutation(mutate))


class SafeguardEvidenceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.deck_evidence = importlib.import_module("deck_evidence")
        self.source = json.loads(SAFEGUARDS.read_text())

    def test_current_safeguard_evidence_is_valid(self) -> None:
        self.assertIs(
            self.deck_evidence.validate_safeguard_evidence(self.source),
            self.source,
        )

    def test_failed_or_inconsistent_verdict_is_rejected(self) -> None:
        for label, mutate in (
            ("case", lambda data: data["cases"][0].__setitem__("passed", False)),
            ("summary", lambda data: data.__setitem__("all_passed", False)),
            ("wrong type", lambda data: data["cases"][0].__setitem__("passed", 1)),
        ):
            with self.subTest(label=label):
                evidence = copy.deepcopy(self.source)
                mutate(evidence)
                with self.assertRaisesRegex(ValueError, "passed|verdict"):
                    self.deck_evidence.validate_safeguard_evidence(evidence)

    def test_scenario_names_are_exact_and_unique(self) -> None:
        mutations = (
            lambda data: data["cases"].pop(),
            lambda data: data["cases"][0].__setitem__("name", "wrong_name"),
            lambda data: data["cases"][0].__setitem__("name", data["cases"][1]["name"]),
        )
        for mutate in mutations:
            with self.subTest(mutate=mutate):
                evidence = copy.deepcopy(self.source)
                mutate(evidence)
                with self.assertRaisesRegex(ValueError, "scenario|case|name|coverage|duplicate"):
                    self.deck_evidence.validate_safeguard_evidence(evidence)

    def test_unhashable_or_non_string_scenario_name_is_clean_value_error(self) -> None:
        for bad_name in ({"not": "hashable"}, ["not", "hashable"], 7, None):
            with self.subTest(bad_name=bad_name):
                evidence = copy.deepcopy(self.source)
                evidence["cases"][0]["name"] = bad_name
                with self.assertRaisesRegex(ValueError, "name.*string"):
                    self.deck_evidence.validate_safeguard_evidence(evidence)

    def test_real_pqc_readiness_is_required(self) -> None:
        mutations = (
            lambda data: data.__setitem__("pqc_ready", False),
            lambda data: data["pqc_profile"].__setitem__("ml_kem_768", False),
            lambda data: data["pqc_profile"].__setitem__("ml_dsa_65", "true"),
        )
        for mutate in mutations:
            with self.subTest(mutate=mutate):
                evidence = copy.deepcopy(self.source)
                mutate(evidence)
                with self.assertRaisesRegex(ValueError, "PQC|ML-KEM|ML-DSA|pqc"):
                    self.deck_evidence.validate_safeguard_evidence(evidence)

    def test_code_hash_drift_and_missing_required_coverage_are_rejected(self) -> None:
        for label, mutate in (
            (
                "drift",
                lambda data: data["code_sha256"].__setitem__(
                    "prototypes/nishan_pq/nishan/core.py", "0" * 64
                ),
            ),
            (
                "missing",
                lambda data: data["code_sha256"].pop(
                    "prototypes/nishan_pq/tests/test_pdf_format_policy.py"
                ),
            ),
        ):
            with self.subTest(label=label):
                evidence = copy.deepcopy(self.source)
                mutate(evidence)
                with self.assertRaisesRegex(ValueError, "hash|SHA-256|coverage"):
                    self.deck_evidence.validate_safeguard_evidence(evidence)

    def test_schema_and_top_level_types_are_strict(self) -> None:
        mutations = (
            lambda data: data.__setitem__("version", "wrong/version"),
            lambda data: data.__setitem__("cases", {}),
            lambda data: data.__setitem__("code_sha256", []),
        )
        for mutate in mutations:
            with self.subTest(mutate=mutate):
                evidence = copy.deepcopy(self.source)
                mutate(evidence)
                with self.assertRaises(ValueError):
                    self.deck_evidence.validate_safeguard_evidence(evidence)

    def test_required_profile_fields_and_types_are_strict(self) -> None:
        mutations = (
            lambda data: data["profile"].pop("roster_size"),
            lambda data: data["profile"].__setitem__("coalition_limit", True),
            lambda data: data["profile"].__setitem__("familywise_epsilon", "1e-6"),
            lambda data: data["profile"].__setitem__("fusion_policy", "legacy"),
            lambda data: data["profile"].__setitem__("fixture", "production corpus"),
            lambda data: data["profile"].__setitem__("pin_provisioning", None),
        )
        for mutate in mutations:
            with self.subTest(mutate=mutate):
                evidence = copy.deepcopy(self.source)
                mutate(evidence)
                with self.assertRaisesRegex(ValueError, "profile"):
                    self.deck_evidence.validate_safeguard_evidence(evidence)

    def test_each_case_observation_enforces_claim_invariants(self) -> None:
        mutations = {
            "concurrent_distinct_sessions_rows": lambda observed: observed.__setitem__("rows", [0, 0]),
            "third_release_row": lambda observed: observed.__setitem__("row", "2"),
            "signed_copy_hash_match": lambda observed: observed[0].__setitem__("actual", "0" * 64),
            "clean_pdf_trace": lambda observed: observed["decision"].__setitem__("decision", "visual_only"),
            "trace_read_only": lambda observed: observed.pop("after_sha256"),
            "rendered_transplant_abstains": lambda observed: observed.__setitem__("attribution", [{}]),
            "low_capacity_pdf_abstains": lambda observed: observed.__setitem__("attribution", None),
            "wrong_pin_rejected": lambda observed: observed.__setitem__("output_exists", True),
            "rollback_release_rejected": lambda observed: observed.__setitem__("comparison", "consistent"),
            "rollback_trace_rejected": lambda observed: observed.__setitem__("rejected", "true"),
            "checkpoint_failure_no_publication": lambda observed: observed.__setitem__("destination_preserved", False),
            "checkpoint_failure_preserves_existing_output": lambda observed: observed.__setitem__("existing_destination", False),
        }
        for case_name, mutate in mutations.items():
            with self.subTest(case_name=case_name):
                evidence = copy.deepcopy(self.source)
                case = next(item for item in evidence["cases"] if item["name"] == case_name)
                mutate(case["observed"])
                with self.assertRaisesRegex(ValueError, case_name):
                    self.deck_evidence.validate_safeguard_evidence(evidence)


class DeckAndVerifierTests(unittest.TestCase):
    maxDiff = None

    @classmethod
    def setUpClass(cls) -> None:
        cls.build_deck = importlib.import_module("build_deck")
        cls.deck_evidence = importlib.import_module("deck_evidence")
        cls.verify_submission = importlib.import_module("verify_submission")

    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.output_dir = self.root / "submissions" / "SIH26237_NISHAN_PQ"
        self.output_dir.mkdir(parents=True)
        self.pptx = self.output_dir / "NISHAN-PQ_SIH26237.pptx"
        self.pdf = self.output_dir / "NISHAN-PQ_SIH26237.pdf"
        self.notes = self.output_dir / "NISHAN-PQ_SIH26237_speaker_notes.md"
        self.team_id = "TEAM-26237"
        self.team_name = "NISHAN TEST TEAM"
        self.repository_url = "https://example.org/nishan/source"

        self.build_deck.build(
            ROOT / "resources" / "SIH2026-IDEA-Presentation-Format.pptx",
            ROOT / "artifacts" / "nishan" / "attribution-scores.png",
            ROOT / "artifacts" / "nishan" / "measured-results.json",
            ROOT / "artifacts" / "nishan" / "tardos-pdf-benchmark.json",
            ROOT / "artifacts" / "nishan" / "tardos-independent-codebook-study.json",
            ROOT / "artifacts" / "nishan" / "end-to-end-jpeg-trials.json",
            self.pptx,
            self.team_id,
            self.team_name,
            physical_path=PHYSICAL,
            repository_url=self.repository_url,
        )

    def prepare_verifier_root(self) -> None:
        artifact_dir = self.root / "artifacts" / "nishan"
        artifact_dir.mkdir(parents=True)
        for name in (
            "measured-results.json",
            "tardos-pdf-benchmark.json",
            "tardos-independent-codebook-study.json",
            "end-to-end-jpeg-trials.json",
        ):
            shutil.copy2(ROOT / "artifacts" / "nishan" / name, artifact_dir / name)
        physical_dir = self.root / "research" / "evidence" / "nishan-bias-physical-2026-09-10"
        physical_dir.mkdir(parents=True)
        shutil.copy2(PHYSICAL, physical_dir / "scores.json")
        safeguard_dir = self.root / "research" / "evidence" / "nishan-selection-2026-09-12" / "demo"
        safeguard_dir.mkdir(parents=True)
        shutil.copy2(SAFEGUARDS, safeguard_dir / "results.json")
        self.render_pdf()

    def render_pdf(self, stale: str = "") -> None:
        prs = Presentation(self.pptx)
        document = fitz.open()
        for index, content in enumerate(slide_text(prs)):
            page = document.new_page(width=960, height=540)
            prefix = stale + "\n" if index == 0 and stale else ""
            page.insert_textbox(fitz.Rect(18, 18, 942, 522), prefix + content, fontsize=5.8)
        document.save(self.pdf)
        document.close()

    def test_build_has_six_editable_slides_and_complete_matching_notes(self) -> None:
        prs = Presentation(self.pptx)
        self.assertEqual(len(prs.slides), 6)
        texts = slide_text(prs)
        combined = "\n".join(texts)
        self.assertIn(self.team_id, combined)
        self.assertIn(self.team_name, combined)
        self.assertIn(self.repository_url, combined)
        self.assertNotRegex(combined, r"\[(?:ENTER|INSERT)[^]]+\]")
        self.assertNotIn("real physical captures pending", combined)
        self.assertNotIn("Synthetic perspective passes; hardware pending", combined)
        embedded = [note_text(slide) for slide in prs.slides]
        self.assertTrue(all(embedded))
        exported = self.notes.read_text()
        for index, block in enumerate(embedded, 1):
            self.assertIn(f"## Slide {index}", exported)
            self.assertIn(block, exported)
        joined = "\n".join(embedded)
        for prefix, count in (("R", 14), ("P", 24)):
            for number in range(1, count + 1):
                self.assertEqual(joined.count(f"{prefix}{number}:"), 1)
        semantic_anchors = {
            "R8:": "NOT MET",
            "R14:": "No public blockchain",
            "P1:": "separate witness is now called",
            "P2:": "corrected",
            "P3:": "renamed/prefixed PDFs",
            "P10:": "same akshay3 capture recovered in both—not two recovered captures",
            "P11:": "one recipient identity across 30 accumulated issued sessions",
            "P23:": "stray file '=4.11'",
            "P24:": "Novelty is unestablished",
        }
        for label, anchor in semantic_anchors.items():
            line = next(item for item in joined.splitlines() if item.startswith(label))
            self.assertIn(anchor.lower(), line.lower())

    def test_slide_claims_are_evidence_derived(self) -> None:
        combined = "\n".join(slide_text(Presentation(self.pptx)))
        self.assertIn("30/30", combined)
        self.assertIn("0/4", combined)
        self.assertIn("1/4", combined)
        self.assertIn("12/12", combined)
        self.assertIn("59/59 final-code tests", combined)
        self.assertIn("19 hashes verified", combined)
        self.assertIn("JPEG-Q55 visual-channel trials", combined)
        self.assertIn("same 1/4 capture", combined)
        self.assertIn("Digital PDF corroboration; raster-only matches are research leads", combined)
        self.assertIn("Portable independent evidence verifier", combined)
        self.assertIn("source-copy evidence—not a claim about which human", combined)
        self.assertIn("checks signed witness state under a caller-provisioned key pin", combined)
        self.assertIn("Wrap key", combined)
        self.assertNotIn("Authorize\nML-KEM-768", combined)
        self.assertIn("Six-person team · existing offline laptop prototype", combined)

    def test_slide_two_headline_clears_official_logo(self) -> None:
        prs = Presentation(self.pptx)
        headlines = [
            shape
            for shape in prs.slides[1].shapes
            if getattr(shape, "text", "").startswith("One encrypted source.")
        ]
        self.assertEqual(len(headlines), 1)
        headline = headlines[0]
        right_edge_inches = (headline.left + headline.width) / Inches(1)
        self.assertLessEqual(right_edge_inches, 10.3)
        self.assertEqual(
            headline.text,
            "One encrypted source. Signed, session-specific releases.",
        )

    def test_build_rejects_safeguard_code_hash_drift(self) -> None:
        evidence = json.loads(SAFEGUARDS.read_text())
        evidence["code_sha256"]["prototypes/nishan_pq/nishan/core.py"] = "0" * 64
        bad_path = self.root / "stale-safeguards.json"
        bad_path.write_text(json.dumps(evidence))
        with self.assertRaisesRegex(ValueError, "SHA-256 drift"):
            self.build_deck.build(
                ROOT / "resources" / "SIH2026-IDEA-Presentation-Format.pptx",
                ROOT / "artifacts" / "nishan" / "attribution-scores.png",
                ROOT / "artifacts" / "nishan" / "measured-results.json",
                ROOT / "artifacts" / "nishan" / "tardos-pdf-benchmark.json",
                ROOT / "artifacts" / "nishan" / "tardos-independent-codebook-study.json",
                ROOT / "artifacts" / "nishan" / "end-to-end-jpeg-trials.json",
                self.root / "must-not-build.pptx",
                self.team_id,
                self.team_name,
                physical_path=PHYSICAL,
                safeguard_path=bad_path,
                repository_url=self.repository_url,
            )

    def test_build_rejects_counterfactual_safeguard_observation(self) -> None:
        evidence = json.loads(SAFEGUARDS.read_text())
        wrong_pin = next(
            item for item in evidence["cases"] if item["name"] == "wrong_pin_rejected"
        )
        wrong_pin["observed"]["output_exists"] = True
        bad_path = self.root / "counterfactual-safeguards.json"
        bad_path.write_text(json.dumps(evidence))
        with self.assertRaisesRegex(ValueError, "wrong_pin_rejected"):
            self.build_deck.build(
                ROOT / "resources" / "SIH2026-IDEA-Presentation-Format.pptx",
                ROOT / "artifacts" / "nishan" / "attribution-scores.png",
                ROOT / "artifacts" / "nishan" / "measured-results.json",
                ROOT / "artifacts" / "nishan" / "tardos-pdf-benchmark.json",
                ROOT / "artifacts" / "nishan" / "tardos-independent-codebook-study.json",
                ROOT / "artifacts" / "nishan" / "end-to-end-jpeg-trials.json",
                self.root / "must-not-build-counterfactual.pptx",
                self.team_id,
                self.team_name,
                physical_path=PHYSICAL,
                safeguard_path=bad_path,
                repository_url=self.repository_url,
            )

    def test_displayed_counts_follow_mutated_json_copies(self) -> None:
        end_to_end = json.loads(
            (ROOT / "artifacts" / "nishan" / "end-to-end-jpeg-trials.json").read_text()
        )
        extra_trial = copy.deepcopy(end_to_end["trials"][-1])
        extra_trial["trial"] = 30
        extra_trial["tardos_user_index"] = 30
        extra_trial["session_id"] = "mutated-test-session-30"
        end_to_end["trials"].append(extra_trial)
        end_to_end["parameters"]["trials"] = 31
        end_to_end["parameters"]["roster_rows_scored_per_trial"] = 999
        end_to_end["summary"]["exact_session_attributions"] = 29
        end_to_end["summary"]["false_accused_codebook_rows_total"] = 2
        end_to_end["summary"]["expected_score_mean"] = 1_234.5
        end_to_end["summary"]["expected_score_population_sd"] = 7.25

        trials = json.loads(
            (ROOT / "artifacts" / "nishan" / "tardos-independent-codebook-study.json").read_text()
        )
        trials["parameters"]["trials"] = 31
        for summary in trials["summary"].values():
            summary["trials"] = 31
            summary["all_colluders_recovered"] = 31

        benchmark = json.loads(
            (ROOT / "artifacts" / "nishan" / "tardos-pdf-benchmark.json").read_text()
        )
        benchmark["carrier"]["metrics_by_user"].append(
            copy.deepcopy(benchmark["carrier"]["metrics_by_user"][-1])
        )

        mutated = self.root / "mutated"
        mutated.mkdir()
        end_to_end_path = mutated / "end-to-end.json"
        trials_path = mutated / "trials.json"
        benchmark_path = mutated / "benchmark.json"
        end_to_end_path.write_text(json.dumps(end_to_end))
        trials_path.write_text(json.dumps(trials))
        benchmark_path.write_text(json.dumps(benchmark))
        output = mutated / "mutated.pptx"

        self.build_deck.build(
            ROOT / "resources" / "SIH2026-IDEA-Presentation-Format.pptx",
            ROOT / "artifacts" / "nishan" / "attribution-scores.png",
            ROOT / "artifacts" / "nishan" / "measured-results.json",
            benchmark_path,
            trials_path,
            end_to_end_path,
            output,
            self.team_id,
            self.team_name,
            physical_path=PHYSICAL,
            repository_url=self.repository_url,
        )
        prs = Presentation(output)
        combined = "\n".join(slide_text(prs))
        notes = "\n".join(note_text(slide) for slide in prs.slides)
        self.assertIn("29/31 JPEG-Q55 visual-channel trials", combined)
        self.assertIn("6 one-page PDFs retained text", combined)
        self.assertIn("R3: Visually identical / forensically distinct — 6 one-page PDFs", notes)
        self.assertIn("P11: The 29/31 end-to-end JPEG-Q55 runs", notes)
        self.assertIn("rows 0..30", notes)
        self.assertIn("other-issued score is null 1 time and non-null 30 times", notes)
        self.assertIn("P12: All 31 end-to-end trials", notes)
        self.assertIn("P13: The 31-codebook × 4-strategy study has 124", notes)
        self.assertIn("P19: Registration and 999-row scoring", notes)

    def test_verifier_rejects_stale_pptx_claim(self) -> None:
        self.prepare_verifier_root()
        prs = Presentation(self.pptx)
        prs.slides[0].shapes.add_textbox(Inches(0.2), Inches(0.2), Inches(5), Inches(0.3)).text = (
            "real physical captures pending"
        )
        prs.save(self.pptx)
        result = self.verify_submission.verify(self.root, allow_placeholders=True)
        self.assertFalse(result["ok"])
        self.assertTrue(any("stale" in error.lower() and "PPTX" in error for error in result["errors"]))

    def test_verifier_accepts_matching_generated_package(self) -> None:
        self.prepare_verifier_root()
        for allow_placeholders in (True, False):
            with self.subTest(allow_placeholders=allow_placeholders):
                result = self.verify_submission.verify(self.root, allow_placeholders)
                self.assertTrue(result["ok"], result["errors"])
                self.assertEqual(result["fresh_safeguards"], {
                    "cases": 12,
                    "all_passed": True,
                    "pqc_ready": True,
                    "verified_code_hashes": 19,
                })

    def test_verifier_rejects_stale_pdf_claim(self) -> None:
        self.prepare_verifier_root()
        self.render_pdf("Synthetic perspective passes; hardware pending")
        result = self.verify_submission.verify(self.root, allow_placeholders=True)
        self.assertFalse(result["ok"])
        self.assertTrue(any("stale" in error.lower() and "PDF" in error for error in result["errors"]))

    def test_verifier_rejects_missing_physical_evidence(self) -> None:
        self.prepare_verifier_root()
        physical = self.root / "research" / "evidence" / "nishan-bias-physical-2026-09-10" / "scores.json"
        physical.unlink()
        result = self.verify_submission.verify(self.root, allow_placeholders=True)
        self.assertFalse(result["ok"])
        self.assertTrue(any("physical evidence" in error.lower() and "missing" in error.lower() for error in result["errors"]))

    def test_verifier_rejects_missing_safeguard_evidence(self) -> None:
        self.prepare_verifier_root()
        safeguard = self.root / "research" / "evidence" / "nishan-selection-2026-09-12" / "demo" / "results.json"
        safeguard.unlink()
        result = self.verify_submission.verify(self.root, allow_placeholders=True)
        self.assertFalse(result["ok"])
        self.assertTrue(any("safeguard evidence" in error.lower() for error in result["errors"]))

    def test_verifier_rejects_failed_safeguard_evidence(self) -> None:
        self.prepare_verifier_root()
        safeguard = self.root / "research" / "evidence" / "nishan-selection-2026-09-12" / "demo" / "results.json"
        evidence = json.loads(safeguard.read_text())
        evidence["cases"][0]["passed"] = False
        safeguard.write_text(json.dumps(evidence))
        result = self.verify_submission.verify(self.root, allow_placeholders=True)
        self.assertFalse(result["ok"])
        self.assertTrue(any("safeguard evidence" in error.lower() for error in result["errors"]))

    def test_verifier_rejects_malformed_safeguard_observation(self) -> None:
        self.prepare_verifier_root()
        safeguard = self.root / "research" / "evidence" / "nishan-selection-2026-09-12" / "demo" / "results.json"
        evidence = json.loads(safeguard.read_text())
        trace = next(item for item in evidence["cases"] if item["name"] == "trace_read_only")
        trace["observed"].pop("after_sha256")
        safeguard.write_text(json.dumps(evidence))
        result = self.verify_submission.verify(self.root, allow_placeholders=True)
        self.assertFalse(result["ok"])
        self.assertTrue(any(
            "safeguard evidence" in error.lower() and "trace_read_only" in error
            for error in result["errors"]
        ))

    def test_handoff_prose_and_runbook_use_current_boundaries(self) -> None:
        submission = (PACKAGE / "paste_ready_submission.md").read_text()
        runbook = (PACKAGE / "DEMO_RUNBOOK.md").read_text()
        combined = submission + "\n" + runbook
        normalized = " ".join(combined.split()).lower()
        for phrase in (
            "All 12 named offline safeguard scenarios passed",
            "59/59 tests",
            "All 19 code hashes",
            "shipped historical threshold recovered 0/4",
            "same one capture",
            "configured release rejected a wrong pin",
            "signed release-authorization record remains",
            "portable independent evidence verifier remain future work",
        ):
            self.assertIn(phrase.lower(), normalized)
        for name in self.deck_evidence.EXPECTED_SAFEGUARD_CASES:
            self.assertIn(name, runbook)
        self.assertEqual(runbook.count("/tmp/nishan-judges-run-01"), 3)
        self.assertNotIn("demo-judges-run-01", runbook.replace("/tmp/nishan-judges-run-01", ""))
        self.assertNotIn("rm ", runbook.lower())
        self.assertNotIn("trash", runbook.lower())
        for stale in (
            "real physical captures pending",
            "witness is not enforced",
            "allocation is not atomic",
            "30/30 attribution success",
        ):
            self.assertNotIn(stale.lower(), normalized)


if __name__ == "__main__":
    unittest.main()
